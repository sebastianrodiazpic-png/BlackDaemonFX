"""Exportador XLSX del estado completo del daemon, con auditoria explicable.

Genera el libro principal del proyecto (por defecto
`reports/deriv_demo_trades.xlsx`) con ocho hojas: Trades, Open Positions,
Summary, Instruments, Account, Audit Log, "Entrada vs Ahora" y Metadata.

Dos ideas de diseno importantes:

1. **Explicabilidad**: la hoja Trades no se limita a los datos crudos; anade
   columnas en espanol que reconstruyen POR QUE se entro (confirmaciones
   cumplidas y faltantes, decision, calidad) y COMO se cerro.
2. **Historial permanente**: si el repositorio expone
   `trade_history_dataframe`, el libro se reconstruye desde el journal
   permanente, de modo que un reinicio de estadisticas no borre el historial
   ya exportado.

Vinculaciones:
    - `database.repository.TradingRepository`: origen de todos los DataFrames.
    - `trade_outcome_policy`: `is_break_even_rr` y `decisive_outcome` deciden
      la clasificacion del cierre, compartida con el resto del proyecto.
    - Consumidor: `reporting.trade_reporting_service`, que lo invoca tras cada
      evento relevante del ciclo de vida.
"""

from __future__ import annotations

import json
import math
import os
import uuid
import inspect
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd

from trade_outcome_policy import decisive_outcome, is_break_even_rr


class TradeReportExporter:
    """Exporta el estado del daemon a XLSX con auditoría explicable.

    Además de las hojas históricas, la hoja Trades incorpora columnas legibles
    en español con las razones de entrada y la clasificación del cierre. La hoja
    Summary incorpora KPIs de ganancia/pérdida y un gráfico circular de TP1,
    TP2, Stop Loss y Break Even/Otros.
    """

    DISPLAY_TIMEZONE = "America/Santiago"
    TELEMETRY_ROWS = 1000
    AUDIT_ROWS = 2000
    AUDIT_EXCEL_SUMMARY_LIMIT = 4000
    AUDIT_EXCEL_SAFE_CELL_LIMIT = 30000

    CONFIRMATION_LABELS = {
        "h1_trend": "Tendencia H1",
        "m15_setup": "Setup M15",
        "liquidity_sweep": "Barrido de liquidez",
        "m15_structure": "Estructura M15 / CHOCH-BOS",
        "premium_discount": "Premium / Discount",
        "fresh_order_block": "Order Block fresco",
        "clean_retest": "Retesteo limpio",
        "rejection": "Rechazo del precio",
        "displacement": "Desplazamiento",
        "micro_structure": "Microestructura M5",
        "strong_close": "Cierre fuerte",
        "m1_follow_through": "Segunda vela M1 de continuación",
        "directional_candle": "Vela direccional",
        "confirmation_mode_ok": "Modo de confirmación",
        "momentum": "Momentum",
        "harmonic_confirmation": "Patrón armónico",
        "h1_extreme_doji_confirmation": "Doji H1 en extremo",
    }

    DECISION_LABELS = {
        "STRICT_CONFIRMED": "CONFIRMADA ESTRICTAMENTE",
        "ADAPTIVE_75_CONFIRMED": "CONFIRMADA POR REGLA ADAPTATIVA >=75%",
        "ADAPTIVE_80_CONFIRMED": "CONFIRMADA POR REGLA ADAPTATIVA >=80%",
        "REJECTED": "RECHAZADA",
    }

    EXECUTION_MODE_LABELS = {
        "SPLIT": "DOS ENTRADAS: 0,5% + 0,5%",
        "SINGLE_FALLBACK": "ENTRADA ÚNICA ALTERNATIVA: HASTA 1%",
        "SINGLE": "ENTRADA ÚNICA",
    }

    def __init__(self, repository, output_path=None):
        """Fija repositorio y destino del libro.

        Si no se indica ruta, usa `reports/deriv_demo_trades.xlsx` relativo al
        directorio de trabajo.
        """
        self.repository = repository
        if output_path is None:
            output_path = Path("reports") / "deriv_demo_trades.xlsx"
        self.output_path = Path(output_path)

    @staticmethod
    def _bounded_read(reader, source, limit):
        if "limit" in inspect.signature(reader).parameters:
            return reader(source=source, limit=limit)
        return reader(source=source).tail(limit)

    def export(self, source="DEMO"):
        """Reconstruye el libro completo desde cero para un entorno dado.

        Args:
            source: DEMO, LIVE o PAPER.

        Returns:
            dict con `path` absoluto, `total_trades`, `open_positions`,
            el `summary` agregado y el numero de snapshots Entrada vs Ahora.

        Pasos: lee el historial (journal permanente si existe, tabla operativa
        si no), convierte marcas de tiempo a hora de Chile, enriquece con las
        columnas explicables, compacta la bitacora de auditoria (sus payloads
        JSON completos NO caben en Excel; quedan en la base) y escribe las ocho
        hojas antes de aplicar el formato.

        Todas las lecturas opcionales se resuelven con `getattr`+`callable`
        para tolerar repositorios reducidos o dobles de test.
        """
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

        # v46: el XLSX se reconstruye desde el journal permanente cuando existe.
        # Así un reset de la tabla operativa no elimina entradas históricas del reporte.
        history_reader = getattr(self.repository, "trade_history_dataframe", None)
        if callable(history_reader):
            trades = history_reader(source=source)
        else:
            trades = self.repository.trades_dataframe(source=source)
        trades = self._localize_dataframe_timestamps(trades)
        trades = self._enrich_trades_for_report(trades)

        summary = self.repository.summary(source=source)
        open_positions = pd.DataFrame()
        if not trades.empty and "status" in trades.columns:
            open_positions = trades[
                trades["status"].astype(str).str.upper() == "OPEN"
            ].copy()

        summary_df = pd.DataFrame(
            [{"metric": key, "value": value} for key, value in summary.items() if key != "instruments"]
        )
        instruments = summary.get("instruments", [])
        instruments_df = pd.DataFrame({"instrument": instruments}) if instruments else pd.DataFrame(columns=["instrument"])

        account_df = self._localize_dataframe_timestamps(self._get_account_snapshots())
        audit_df = pd.DataFrame()
        audit_reader = getattr(self.repository, "audit_events_dataframe", None)
        if callable(audit_reader):
            audit_df = self._localize_dataframe_timestamps(self._bounded_read(audit_reader, source, self.AUDIT_ROWS))
            audit_df = self._compact_audit_dataframe(audit_df)

        # v68: dataset longitudinal Entrada vs. Ahora para investigación.
        entry_vs_now_df = pd.DataFrame()
        entry_now_reader = getattr(self.repository, "trade_audit_snapshots_dataframe", None)
        if callable(entry_now_reader):
            entry_vs_now_df = self._localize_dataframe_timestamps(self._bounded_read(entry_now_reader, source, self.TELEMETRY_ROWS))
            if not entry_vs_now_df.empty:
                for col in ("entry_view_json", "current_view_json", "market_json", "visual_context_json"):
                    if col in entry_vs_now_df.columns:
                        entry_vs_now_df[col] = entry_vs_now_df[col].map(
                            lambda value: str(value)[: 1000] if value is not None else ""
                        )
        metadata_df = pd.DataFrame([
            {
                "source": str(source).upper(),
                "generated_at_utc": datetime.now(timezone.utc).isoformat(),
                "generated_at_chile": self._format_chile_now(),
                "display_timezone": self.DISPLAY_TIMEZONE,
                "total_trades": int(summary.get("total_trades", 0)),
                "open_trades": int(summary.get("open_trades", 0)),
                "closed_trades": int(summary.get("closed_trades", 0)),
                "report_version": "bounded-telemetry-v2",
                "audit_row_limit": self.AUDIT_ROWS,
                "snapshot_row_limit": self.TELEMETRY_ROWS,
                "full_history": "SQLite; Excel contains recent telemetry only",
                "entry_vs_now_snapshots": int(len(entry_vs_now_df)),
            }
        ])

        temporary = self.output_path.with_name(self.output_path.stem + "." + uuid.uuid4().hex + ".tmp.xlsx")
        try:
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                trades.to_excel(writer, sheet_name="Trades", index=False)
                open_positions.to_excel(writer, sheet_name="Open Positions", index=False)
                summary_df.to_excel(writer, sheet_name="Summary", index=False)
                instruments_df.to_excel(writer, sheet_name="Instruments", index=False)
                account_df.to_excel(writer, sheet_name="Account", index=False)
                audit_df.to_excel(writer, sheet_name="Audit Log", index=False)
                entry_vs_now_df.to_excel(writer, sheet_name="Entrada vs Ahora", index=False)
                metadata_df.to_excel(writer, sheet_name="Metadata", index=False)

            self._format_workbook(path=temporary)
            os.replace(temporary, self.output_path)
        finally:
            temporary.unlink(missing_ok=True)

        return {
            "path": str(self.output_path.resolve()),
            "total_trades": len(trades),
            "open_positions": len(open_positions),
            "summary": summary,
            "entry_vs_now_snapshots": len(entry_vs_now_df),
        }

    # ------------------------------------------------------------------
    # AUDIT LOG COMPACTO PARA EXCEL
    # ------------------------------------------------------------------

    @classmethod
    def _audit_payload_dict(cls, value):
        """Normaliza un payload de auditoria a diccionario.

        Acepta dict, JSON en texto o escalares. Lo que no sea un objeto se
        envuelve en `{"value": ...}` y un JSON corrupto se conserva como texto
        en vez de perderse.
        """
        if isinstance(value, dict):
            return value
        if isinstance(value, str) and value.strip():
            try:
                parsed = json.loads(value)
                return parsed if isinstance(parsed, dict) else {"value": parsed}
            except Exception:
                return {"value": value}
        if value is None:
            return {}
        return {"value": value}

    @classmethod
    def _audit_summary(cls, payload):
        """Construye un resumen útil sin exportar el JSON completo a una celda."""
        obj = cls._audit_payload_dict(payload)
        parts = []

        for key, label in (
            ("bot_profile", "Bot"),
            ("daemon_magic", "Magic"),
            ("phase", "Fase"),
            ("elapsed_seconds", "Duración"),
            ("index", "Índice"),
            ("total", "Total"),
        ):
            value = obj.get(key)
            if value not in (None, "", [], {}):
                parts.append(f"{label}: {value}")

        result = obj.get("result")
        if isinstance(result, dict):
            for key, label in (
                ("action", "Acción"),
                ("reason", "Motivo"),
                ("error", "Error"),
                ("strategy_name", "Estrategia"),
                ("execution_mode", "Ejecución"),
            ):
                value = result.get(key)
                if value not in (None, "", [], {}):
                    parts.append(f"{label}: {value}")

            confirmations = result.get("confirmation_percentage")
            if confirmations not in (None, ""):
                parts.append(f"Confirmaciones: {confirmations}")

        recovery = obj.get("persistence_recovery")
        if isinstance(recovery, dict) and recovery:
            imported = recovery.get("imported_daemon")
            if imported:
                parts.append(f"Recuperadas MT5→DB: {imported}")

        text = " | ".join(str(part) for part in parts)
        if not text:
            # Un fallback corto es útil para eventos no estructurados.
            text = json.dumps(obj, ensure_ascii=False, default=str)

        limit = int(cls.AUDIT_EXCEL_SUMMARY_LIMIT)
        if len(text) > limit:
            text = text[: max(0, limit - 24)] + "… [RESUMEN TRUNCADO]"
        return text

    @classmethod
    def _compact_audit_dataframe(cls, dataframe: pd.DataFrame) -> pd.DataFrame:
        """Reduce la bitacora a columnas fijas mas un `resumen` de texto.

        Los payloads completos pueden ser enormes y superar el limite de celda
        de Excel, asi que se sustituyen por un resumen y se marca
        `payload_completo_en_db = SI`, dejando claro que el dato integro sigue
        disponible en SQLite.

        Devuelve un DataFrame vacio pero CON las columnas esperadas si no hay
        eventos, para que la hoja mantenga su cabecera.
        """
        if dataframe is None or dataframe.empty:
            return pd.DataFrame(columns=[
                "id", "event_time", "source", "bot_profile", "daemon_magic",
                "event_type", "instrument", "action", "reason",
                "execution_key", "broker_position_ticket", "cycle_number",
                "resumen", "payload_completo_en_db",
            ])

        df = dataframe.copy()
        payload_series = (
            df["payload"]
            if "payload" in df.columns
            else df["payload_json"] if "payload_json" in df.columns
            else pd.Series([{}] * len(df), index=df.index)
        )

        payload_objects = payload_series.apply(cls._audit_payload_dict)
        df["bot_profile"] = payload_objects.apply(
            lambda obj: obj.get("bot_profile") if isinstance(obj, dict) else None
        )
        df["daemon_magic"] = payload_objects.apply(
            lambda obj: obj.get("daemon_magic") if isinstance(obj, dict) else None
        )
        df["resumen"] = payload_objects.apply(cls._audit_summary)
        df["payload_completo_en_db"] = "SÍ"

        wanted = [
            "id", "event_time", "source", "bot_profile", "daemon_magic",
            "event_type", "instrument", "action", "reason",
            "execution_key", "broker_position_ticket", "cycle_number",
            "resumen", "payload_completo_en_db",
        ]
        for column in wanted:
            if column not in df.columns:
                df[column] = None

        # Defensa adicional: ninguna celda textual debe acercarse al límite Excel.
        safe_limit = int(cls.AUDIT_EXCEL_SAFE_CELL_LIMIT)
        for column in df.columns:
            if df[column].dtype == object:
                df[column] = df[column].apply(
                    lambda value: (
                        value[:safe_limit] + "…"
                        if isinstance(value, str) and len(value) > safe_limit
                        else value
                    )
                )

        return df[wanted]

    # ------------------------------------------------------------------
    # ENRIQUECIMIENTO EXPLICABLE DE TRADES
    # ------------------------------------------------------------------

    def _enrich_trades_for_report(self, dataframe: pd.DataFrame) -> pd.DataFrame:
        """Anade a cada trade las columnas explicables en espanol.

        Recorre fila a fila extrayendo los metadatos del JSON `details` y
        construyendo el registro de auditoria legible. Si el DataFrame llega
        vacio, crea igualmente las columnas de `_report_columns()` para que la
        hoja conserve su estructura.
        """
        if dataframe is None:
            return pd.DataFrame()
        df = dataframe.copy()
        if df.empty:
            for name in self._report_columns():
                if name not in df.columns:
                    df[name] = pd.Series(dtype="object")
            return df

        records = []
        for _, row in df.iterrows():
            metadata = self._extract_metadata(row)
            records.append(self._report_record(row, metadata))

        audit = pd.DataFrame(records, index=df.index)
        for column in audit.columns:
            df[column] = audit[column]
        return df

    @classmethod
    def _report_columns(cls):
        """Nombres de las columnas explicables anadidas a la hoja Trades.

        Es la lista canonica: cualquier columna nueva debe declararse aqui para
        que aparezca tambien cuando no hay operaciones.
        """
        return [
            "estrategia_id",
            "version_estrategia",
            "perfil_bot",
            "magic_estrategia",
            "motivo_entrada_es",
            "decision_entrada_es",
            "porcentaje_confirmaciones",
            "confirmaciones_aprobadas",
            "confirmaciones_totales",
            "score_calidad",
            "grado_calidad",
            "confirmaciones_cumplidas_es",
            "confirmaciones_faltantes_es",
            "condiciones_criticas_faltantes_es",
            "modalidad_ejecucion_es",
            "pierna_operacion",
            "confirmacion_armonica_es",
            "divergencia_es",
            "doji_h1_extremo_es",
            "mfe_max_rr",
            "mae_max_rr",
            "etapa_runner",
            "profit_lock_runner_rr",
            "motivo_salida_runner",
            "clasificacion_cierre_es",
            "resultado_monetario_es",
        ]

    def _report_record(self, row, metadata):
        """Construye el diccionario explicable de UNA operacion.

        Args:
            row: fila del DataFrame de trades.
            metadata: `details["metadata"]` ya extraido.

        Returns:
            dict con exactamente las claves de `_report_columns()`.

        Los trades antiguos que no guardaban decision se etiquetan
        explicitamente como "NO REGISTRADA (TRADE LEGACY)" en lugar de quedar
        en blanco, para no confundir un dato ausente con un dato negativo.
        """
        strategy_name = self._clean(metadata.get("strategy_name")) or "SMC"
        strategy_version = (
            self._clean(metadata.get("strategy_version"))
            or self._clean(row.get("strategy_version") if hasattr(row, "get") else None)
        )
        passed = self._as_list(metadata.get("passed_confirmations"))
        missing = self._as_list(metadata.get("missing_confirmations"))
        critical_missing = self._as_list(metadata.get("critical_confirmation_failures"))
        percentage = self._safe_number(metadata.get("confirmation_percentage"))
        passed_count = metadata.get("confirmations_passed")
        total_count = metadata.get("confirmations_total")
        score = self._safe_number(metadata.get("trade_score"))
        grade = self._clean(metadata.get("trade_grade"))
        decision_raw = self._clean(metadata.get("confirmation_decision"))
        decision_es = self.DECISION_LABELS.get(decision_raw, decision_raw or "NO REGISTRADA (TRADE LEGACY)")

        harmonic = bool(metadata.get("harmonic_confirmed", False))
        harmonic_pattern = self._clean(metadata.get("harmonic_pattern"))
        harmonic_score = self._safe_number(metadata.get("harmonic_score"))
        harmonic_es = "Sí"
        if harmonic:
            if harmonic_pattern:
                harmonic_es += f" - {harmonic_pattern}"
            if harmonic_score is not None:
                harmonic_es += f" (score {harmonic_score:.1f})"
        else:
            harmonic_es = "No / no requerida"

        divergence = bool(metadata.get("divergence_confirmation", False))
        divergence_type = self._clean(metadata.get("divergence_type"))
        divergence_es = "Sí" if divergence else "No / no requerida"
        if divergence and divergence_type:
            divergence_es += " - " + divergence_type.replace("_", " ")

        doji = bool(metadata.get("h1_doji_confirmation", False))
        doji_type = self._clean(metadata.get("h1_doji_type"))
        doji_zone = self._clean(metadata.get("h1_doji_zone"))
        doji_es = "Sí" if doji else "No / no requerida"
        if doji and doji_type:
            doji_es += " - " + doji_type.replace("_", " ")
        if doji and doji_zone:
            doji_es += f" ({doji_zone})"

        execution_mode = self._clean(metadata.get("execution_mode"))
        execution_mode_es = self.EXECUTION_MODE_LABELS.get(execution_mode, execution_mode or "NO REGISTRADA (TRADE LEGACY)")
        leg = self._clean(metadata.get("trade_leg")) or "NO REGISTRADA (TRADE LEGACY)"

        motivo = self._build_entry_reason(
            metadata, decision_es, passed, missing, percentage, score,
            harmonic_es, divergence_es, doji_es,
            strategy_name=strategy_name,
        )
        close_class = self._classify_close(row)
        pnl = self._safe_number(row.get("net_pnl")) if hasattr(row, "get") else None
        outcome = decisive_outcome(
            realized_rr=row.get("realized_rr") if hasattr(row, "get") else None,
            net_pnl=pnl,
            status=row.get("status") if hasattr(row, "get") else None,
            classification=close_class,
        )
        monetary = (
            "GANADOR" if outcome == "WIN"
            else "PERDEDOR" if outcome == "LOSS"
            else "BREAK EVEN" if outcome == "BREAK_EVEN"
            else "SIN RESULTADO"
        )

        mfe_rr = self._safe_number(metadata.get("max_favorable_excursion_rr"))
        mae_rr = self._safe_number(metadata.get("max_adverse_excursion_rr"))
        runner_stage = self._clean(metadata.get("runner_extension_stage"))
        runner_lock = self._safe_number(metadata.get("runner_profit_lock_rr"))
        runner_exit_reason = self._clean(metadata.get("runner_exit_reason"))

        return {
            "estrategia_id": strategy_name,
            "version_estrategia": strategy_version,
            "perfil_bot": self._clean(metadata.get("bot_profile")),
            "magic_estrategia": metadata.get("daemon_magic"),
            "motivo_entrada_es": motivo,
            "decision_entrada_es": decision_es,
            "porcentaje_confirmaciones": percentage,
            "confirmaciones_aprobadas": passed_count,
            "confirmaciones_totales": total_count,
            "score_calidad": score,
            "grado_calidad": grade,
            "confirmaciones_cumplidas_es": self._labels(passed),
            "confirmaciones_faltantes_es": self._labels(missing),
            "condiciones_criticas_faltantes_es": self._labels(critical_missing),
            "modalidad_ejecucion_es": execution_mode_es,
            "pierna_operacion": leg,
            "confirmacion_armonica_es": harmonic_es,
            "divergencia_es": divergence_es,
            "doji_h1_extremo_es": doji_es,
            "mfe_max_rr": mfe_rr,
            "mae_max_rr": mae_rr,
            "etapa_runner": runner_stage,
            "profit_lock_runner_rr": runner_lock,
            "motivo_salida_runner": runner_exit_reason,
            "clasificacion_cierre_es": close_class,
            "resultado_monetario_es": monetary,
        }

    def _build_entry_reason(self, metadata, decision_es, passed, missing, percentage, score, harmonic_es, divergence_es, doji_es, strategy_name="SMC"):
        """Redacta en una frase por que se tomo la entrada.

        Encadena estrategia (SMC u ORB), estructura, tendencia H1, zona,
        porcentaje de confirmaciones, calidad, confirmaciones cumplidas y
        faltantes, confluencias (divergencia, armonico, doji H1) y la decision
        final. Solo incluye los fragmentos con dato, de modo que la frase nunca
        contiene huecos.
        """
        if str(strategy_name).upper() == "ORB_NEW_YORK":
            parts = ["Opening Range Breakout New York (ORB)"]
        else:
            parts = ["Smart Money Concepts (SMC)"]
        structure = self._clean(metadata.get("m15_structure_break_type"))
        if structure:
            parts.append(f"estructura {structure}")
        trend = self._clean(metadata.get("h1_trend"))
        if trend:
            parts.append(f"tendencia H1 {trend}")
        zone = self._clean(metadata.get("m15_zone"))
        if zone:
            parts.append(f"zona {zone}")
        if percentage is not None:
            parts.append(f"{percentage:.1f}% de confirmaciones")
        if score is not None:
            parts.append(f"score {score:.1f}/100")
        if passed:
            parts.append("cumplió: " + self._labels(passed))
        if missing:
            parts.append("faltantes no bloqueantes: " + self._labels(missing))
        if divergence_es.startswith("Sí"):
            parts.append("divergencia confirmada")
        if harmonic_es.startswith("Sí"):
            parts.append("confluencia armónica confirmada")
        if doji_es.startswith("Sí"):
            parts.append("Doji H1 confirmado en zona extrema")
        parts.append(f"decisión: {decision_es}")
        return "; ".join(parts) + "."

    def _classify_close(self, row):
        """Clasifica el cierre en TP1..TP4, STOP LOSS, break-even u otro.

        Orden de decision:
        1. Posicion abierta -> "ABIERTA".
        2. Cierre de emergencia por riesgo -> se aparta como "NO CONTABILIZAR",
           porque no refleja la calidad de la estrategia sino una intervencion
           de proteccion.
        3. `is_break_even_rr` (v48): un RR minusculo como +0.01R o -0.01R es
           break-even aunque el PnL monetario tenga signo; asi las comisiones
           no convierten un empate en ganancia o perdida ficticia.
        4. Resto: se deduce el objetivo alcanzado combinando PnL, RR realizado,
           RR planificado y la pierna (TP1 / RUNNER / SINGLE).

        Returns:
            La etiqueta en espanol que alimenta el grafico circular del
            dashboard.
        """
        status = self._clean(row.get("status")).upper()
        if status == "OPEN":
            return "ABIERTA"
        result = self._clean(row.get("result")).upper()
        exit_reason = self._clean(row.get("exit_reason")).upper()
        if "EMERGENCY" in result or "EMERGENCY" in exit_reason or "RIESGO" in result:
            return "EMERGENCIA / NO CONTABILIZAR"

        rr = self._safe_number(row.get("realized_rr"))
        pnl = self._safe_number(row.get("net_pnl"))
        planned_rr = self._safe_number(row.get("planned_rr"))
        leg = self._clean(row.get("pierna_operacion")).upper()

        # v48: +0.01R, -0.01R, etc. son Break Even aunque el PnL tenga signo.
        if is_break_even_rr(rr):
            return "PUNTO DE EQUILIBRIO / OTRO"

        if pnl is not None and pnl > 0:
            if leg == "TP1" or (planned_rr is not None and planned_rr <= 1.25):
                return "TP1"
            if leg in {"RUNNER", "SINGLE", "FULL"} or (planned_rr is not None and planned_rr >= 1.5):
                if rr is not None and rr >= 3.75:
                    return "TP4"
                if rr is not None and rr >= 2.75:
                    return "TP3"
                if rr is None or rr >= 1.5:
                    return "TP2"
                return "GANANCIA PARCIAL / OTRO"
        if rr is not None:
            if rr >= 3.75:
                return "TP4"
            if rr >= 2.75:
                return "TP3"
            if rr >= 1.5:
                return "TP2"
            if 0.75 <= rr < 1.5:
                return "TP1"
            if rr <= -0.75:
                return "STOP LOSS"
        if pnl is not None and pnl < 0:
            return "PÉRDIDA PARCIAL / OTRO"
        if pnl is not None and pnl > 0:
            return "GANANCIA PARCIAL / OTRO"
        return "OTRO"

    def _extract_metadata(self, row):
        """Recupera `details["metadata"]` probando `details_json` y `details`.

        Segun el origen del DataFrame (journal o tabla operativa) el JSON viaja
        en una columna u otra. Devuelve `{}` si no hay nada utilizable.
        """
        candidates = []
        if hasattr(row, "get"):
            candidates.extend([row.get("details_json"), row.get("details")])
        for raw in candidates:
            obj = self._parse_details(raw)
            if isinstance(obj, dict):
                md = obj.get("metadata")
                if isinstance(md, dict):
                    return md
        return {}

    @staticmethod
    def _parse_details(raw):
        """Parsea el JSON `details` devolviendo `{}` ante cualquier problema.

        Un `details` corrupto no debe impedir que el resto del informe se
        exporte.
        """
        if isinstance(raw, dict):
            return raw
        if not isinstance(raw, str) or not raw.strip():
            return {}
        try:
            return json.loads(raw)
        except Exception:
            return {}

    @classmethod
    def _labels(cls, keys):
        """Traduce una lista de confirmaciones a texto separado por comas.

        Devuelve "Ninguna" con lista vacia, para distinguir de forma explicita
        "no falto nada" de una celda en blanco.
        """
        if not keys:
            return "Ninguna"
        return ", ".join(cls.CONFIRMATION_LABELS.get(str(k), str(k).replace("_", " ")) for k in keys)

    @staticmethod
    def _as_list(value):
        """Normaliza a lista: `None` -> `[]`, escalar -> `[escalar]`."""
        if value is None:
            return []
        if isinstance(value, list):
            return value
        if isinstance(value, (tuple, set)):
            return list(value)
        return [value]

    @staticmethod
    def _clean(value):
        """Convierte a texto tratando `None` y `NaN` como cadena vacia.

        El `NaN` de pandas es imprescindible controlarlo: sin esto apareceria
        el literal "nan" en las celdas del informe.
        """
        if value is None:
            return ""
        if isinstance(value, float) and math.isnan(value):
            return ""
        return str(value)

    @staticmethod
    def _safe_number(value):
        """Convierte a float devolviendo `None` ante fallo o `NaN`.

        Devolver `None` en lugar de 0.0 es intencionado: permite distinguir
        "no hay dato" de "el valor es cero".
        """
        if value is None:
            return None
        try:
            number = float(value)
            return None if math.isnan(number) else number
        except (TypeError, ValueError):
            return None

    # ------------------------------------------------------------------
    # ZONA HORARIA CHILE
    # ------------------------------------------------------------------

    @classmethod
    def _format_chile_now(cls) -> str:
        """Momento actual formateado en hora de Chile."""
        local = datetime.now(timezone.utc).astimezone(ZoneInfo(cls.DISPLAY_TIMEZONE))
        return cls._format_local(local)

    @staticmethod
    def _format_local(local: datetime) -> str:
        """Formatea una fecha local anadiendo abreviatura y desfase UTC.

        Deduce CLST (UTC-3, horario de verano) o CLT (UTC-4) a partir del
        desfase real, de modo que el informe no sea ambiguo en los meses de
        cambio horario.
        """
        offset = local.utcoffset()
        hours = int(offset.total_seconds() // 3600) if offset else 0
        abbreviation = "CLST" if hours == -3 else "CLT" if hours == -4 else "CHILE"
        offset_text = local.strftime("%z")
        return local.strftime("%Y-%m-%d %H:%M:%S") + f" {abbreviation} (UTC{offset_text[:3]}:{offset_text[3:]})"

    @classmethod
    def _format_chile_datetime(cls, value):
        """Convierte cualquier representacion de fecha a texto en hora chilena.

        Los valores sin zona horaria se asumen UTC, que es como el proyecto
        persiste todo. Ante un valor no interpretable devuelve el original en
        lugar de descartarlo.
        """
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return value
        try:
            if isinstance(value, datetime):
                dt = value
            else:
                dt = pd.to_datetime(value, errors="coerce", utc=True)
                if pd.isna(dt):
                    return value
                dt = dt.to_pydatetime()
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            else:
                dt = dt.astimezone(timezone.utc)
            return cls._format_local(dt.astimezone(ZoneInfo(cls.DISPLAY_TIMEZONE)))
        except Exception:
            return value

    @classmethod
    def _localize_dataframe_timestamps(cls, dataframe):
        """Pasa a hora de Chile todas las columnas temporales de un DataFrame.

        Detecta las columnas por nombre exacto (`entry_time`, `exit_time`...) o
        porque contengan `_time`/`timestamp`, de forma que las tablas nuevas
        queden cubiertas sin tocar este metodo.
        """
        if dataframe is None or dataframe.empty:
            return dataframe
        df = dataframe.copy()
        timestamp_names = {
            "entry_time", "exit_time", "signal_time", "created_at", "updated_at",
            "snapshot_time", "execution_time", "generated_at", "generated_at_utc",
        }
        for column in list(df.columns):
            if column in timestamp_names or any(token in str(column).lower() for token in ("_time", "timestamp")):
                df[column] = df[column].map(cls._format_chile_datetime)
        return df

    def _get_account_snapshots(self):
        """Snapshots de cuenta, o DataFrame vacio si el repositorio no los expone."""
        if hasattr(self.repository, "account_snapshots_dataframe"):
            reader = self.repository.account_snapshots_dataframe
            return reader(limit=self.TELEMETRY_ROWS) if "limit" in inspect.signature(reader).parameters else reader().tail(self.TELEMETRY_ROWS)
        return pd.DataFrame()

    # ------------------------------------------------------------------
    # FORMATO Y DASHBOARD
    # ------------------------------------------------------------------

    def _format_workbook(self, path=None):
        """Da formato al libro y construye el dashboard de la hoja Summary.

        Si `openpyxl` no esta instalado sale en silencio: el libro ya contiene
        los datos y el formato es puramente cosmetico.
        """
        try:
            from openpyxl import load_workbook
            from openpyxl.chart import PieChart, Reference
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return

        workbook = load_workbook(path or self.output_path)
        blue = "1F4E78"
        white = "FFFFFF"

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color=white)
                cell.fill = PatternFill("solid", fgColor=blue)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for column_cells in worksheet.columns:
                max_length = 0
                for cell in column_cells:
                    value = cell.value
                    if value is not None:
                        max_length = max(max_length, len(str(value)))
                letter = column_cells[0].column_letter
                worksheet.column_dimensions[letter].width = min(max_length + 2, 48)

        trades_ws = workbook["Trades"]
        headers = {cell.value: cell.column for cell in trades_ws[1]}
        for name in (
            "motivo_entrada_es", "confirmaciones_cumplidas_es",
            "confirmaciones_faltantes_es", "condiciones_criticas_faltantes_es",
        ):
            col = headers.get(name)
            if col:
                trades_ws.column_dimensions[trades_ws.cell(1, col).column_letter].width = 48
                for row in range(2, trades_ws.max_row + 1):
                    trades_ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
        pct_col = headers.get("porcentaje_confirmaciones")
        if pct_col:
            for row in range(2, trades_ws.max_row + 1):
                trades_ws.cell(row, pct_col).number_format = '0.00"%"'

        if "Entrada vs Ahora" in workbook.sheetnames:
            evn_ws = workbook["Entrada vs Ahora"]
            evn_headers = {cell.value: cell.column for cell in evn_ws[1]}
            for name in ("entry_view_json", "current_view_json", "market_json", "visual_context_json"):
                col = evn_headers.get(name)
                if col:
                    evn_ws.column_dimensions[evn_ws.cell(1, col).column_letter].width = 48
                    for row in range(2, evn_ws.max_row + 1):
                        evn_ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
            for name in ("entry_score", "entry_confirmation_pct", "current_score", "current_confirmation_pct", "current_rr"):
                col = evn_headers.get(name)
                if col:
                    for row in range(2, evn_ws.max_row + 1):
                        evn_ws.cell(row, col).number_format = "0.00"

        summary_ws = workbook["Summary"]
        self._build_summary_dashboard(summary_ws, trades_ws, PieChart, Reference, Font, PatternFill, Alignment)
        workbook.save(path or self.output_path)
        workbook.close()

    def _build_summary_dashboard(self, ws, trades_ws, PieChart, Reference, Font, PatternFill, Alignment):
        """Dibuja el panel de KPIs y el grafico circular sobre la hoja Summary.

        Los tipos de `openpyxl` llegan por parametro porque el import es
        diferido en `_format_workbook`.

        Es idempotente: limpia primero su area (filas 1-20, columnas D-R) y
        vacia `ws._charts`, de modo que las exportaciones automaticas repetidas
        no acumulen graficos superpuestos.
        """
        blue = "1F4E78"
        white = "FFFFFF"
        header_fill = PatternFill("solid", fgColor=blue)

        # Limpia el área del dashboard para que cada auto-export sea idempotente.
        for row in ws.iter_rows(min_row=1, max_row=20, min_col=4, max_col=18):
            for cell in row:
                cell.value = None
        ws._charts = []

        ws.merge_cells("D1:H1")
        ws["D1"] = "Resumen visual de operaciones"
        ws["D1"].font = Font(bold=True, color=white, size=14)
        ws["D1"].fill = header_fill
        ws["D1"].alignment = Alignment(horizontal="center")

        ws["D3"], ws["E3"] = "Indicador", "Valor"
        for c in (ws["D3"], ws["E3"], ws["G3"], ws["H3"]):
            c.font = Font(bold=True, color=white)
            c.fill = header_fill
        ws["G3"], ws["H3"] = "Tipo de cierre", "Cantidad"

        metrics = [
            ("Trades ganadores", '=COUNTIF(Trades!$*:$*,"GANADOR")'),
        ]
        # Se construyen fórmulas usando las columnas reales del reporte.
        headers = {cell.value: cell.column_letter for cell in trades_ws[1]}
        result_col = headers.get("resultado_monetario_es")
        close_col = headers.get("clasificacion_cierre_es")
        pnl_col = headers.get("net_pnl")
        if not (result_col and close_col and pnl_col):
            return

        # El daemon calcula los KPIs directamente en cada exportación. Así el
        # reporte queda correcto incluso antes de que Excel recalcule fórmulas.
        result_idx = trades_ws[result_col + "1"].column
        close_idx = trades_ws[close_col + "1"].column
        pnl_idx = trades_ws[pnl_col + "1"].column
        winners = losers = 0
        gross_win = gross_loss = 0.0
        close_counts = {
            "TP1": 0,
            "TP2": 0,
            "STOP LOSS": 0,
            "PUNTO DE EQUILIBRIO / OTRO": 0,
            "EMERGENCIA / NO CONTABILIZAR": 0,
        }
        for row_number in range(2, trades_ws.max_row + 1):
            monetary = str(trades_ws.cell(row_number, result_idx).value or "").upper()
            close_type = str(trades_ws.cell(row_number, close_idx).value or "").upper()
            try:
                pnl = float(trades_ws.cell(row_number, pnl_idx).value or 0.0)
            except (TypeError, ValueError):
                pnl = 0.0
            if monetary == "GANADOR":
                winners += 1
                gross_win += pnl
            elif monetary == "PERDEDOR":
                losers += 1
                gross_loss += pnl
            if close_type in close_counts:
                close_counts[close_type] += 1

        total_decided = winners + losers
        win_rate = winners / total_decided if total_decided else 0.0
        loss_rate = losers / total_decided if total_decided else 0.0

        ws["D4"], ws["E4"] = "Trades ganadores", winners
        ws["D5"], ws["E5"] = "Trades perdedores", losers
        ws["D6"], ws["E6"] = "% de trades ganados", win_rate
        ws["D7"], ws["E7"] = "% de trades perdidos", loss_rate
        ws["D8"], ws["E8"] = "Ganancia neta total", gross_win
        ws["D9"], ws["E9"] = "Pérdida neta total", gross_loss
        ws["E6"].number_format = "0.00%"
        ws["E7"].number_format = "0.00%"
        ws["E8"].number_format = '$#,##0.00;[Red]-$#,##0.00'
        ws["E9"].number_format = '$#,##0.00;[Red]-$#,##0.00'

        category_rows = [
            ("TP1", "TP1"),
            ("TP2", "TP2"),
            ("Stop Loss", "STOP LOSS"),
            ("Punto de equilibrio / Otros", "PUNTO DE EQUILIBRIO / OTRO"),
        ]
        for idx, (label, key) in enumerate(category_rows, start=4):
            ws.cell(idx, 7).value = label
            ws.cell(idx, 8).value = close_counts[key]
        ws["G8"] = "Emergencia / No contabilizar"
        ws["H8"] = close_counts["EMERGENCIA / NO CONTABILIZAR"]

        chart = PieChart()
        labels_ref = Reference(ws, min_col=7, min_row=4, max_row=7)
        data_ref = Reference(ws, min_col=8, min_row=3, max_row=7)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(labels_ref)
        chart.title = "Cierres: TP1, TP2, Stop Loss y BE/Otros"
        chart.height = 8.5
        chart.width = 13
        ws.add_chart(chart, "J2")

        ws["D11"] = "Interpretación"
        ws["D11"].font = Font(bold=True, color=white)
        ws["D11"].fill = header_fill
        ws.merge_cells("D11:H11")
        notes = [
            ("TP1", "Primera pierna cerrada aproximadamente en 1R."),
            ("TP2", "Runner/entrada única cerrada aproximadamente en 2R."),
            ("Stop Loss", "Cierre cercano a -1R."),
            ("Ganados/Perdidos", "Se calculan por PnL neto positivo/negativo; emergencias sin PnL no alteran el porcentaje."),
        ]
        for i, (label, text) in enumerate(notes, start=12):
            ws.cell(i, 4).value = label
            ws.cell(i, 5).value = text
            ws.merge_cells(start_row=i, start_column=5, end_row=i, end_column=8)
            ws.cell(i, 5).alignment = Alignment(wrap_text=True)

        for col, width in {"D":24, "E":24, "G":30, "H":14}.items():
            ws.column_dimensions[col].width = width
