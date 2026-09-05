"""Exportador XLSX de la auditoria forense "Entrada vs. Ahora" de UN trade.

A diferencia de `trade_report_exporter` (que exporta la cartera entera), aqui
se disecciona una sola operacion: que se veia al entrar, que se ve ahora y como
evoluciono snapshot a snapshot. Es la herramienta para responder "por que se
cerro esto".

Hojas generadas: Resumen, Tesis Entrada, Ultimo Estado, Mercado Final y
Timeline Completo (una fila por snapshot).

Control de integridad: si la identidad del trade no es verificable, falta la
tesis de entrada o hay snapshots de otro instrumento mezclados, `export` LANZA
`ValueError` en lugar de producir un fichero enganoso. Un informe forense
incorrecto es peor que no tener informe.

Vinculaciones:
    - Payload identico al que consume la pestana web de auditoria, producido por
      `database.repository.TradingRepository.trade_audit_snapshots`.
    - `openpyxl` para el formato final del libro.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd


class TradeAuditExcelExporter:
    """Exporta la auditoría completa Entrada vs. Ahora de un trade a XLSX.

    El payload se obtiene del mismo endpoint SQLAlchemy usado por la pestaña web.
    El archivo está diseñado para análisis posterior: una fila por snapshot en
    Timeline y hojas separadas para tesis de entrada / último estado.
    """

    DISPLAY_TIMEZONE = "America/Santiago"
    SAFE_CELL_LIMIT = 30000

    def __init__(self, output_path):
        """Fija la ruta de destino del XLSX (aun no se crea el fichero)."""
        self.output_path = Path(output_path)

    @classmethod
    def _safe_text(cls, value):
        """Convierte cualquier valor en texto apto para una celda de Excel.

        Estructuras anidadas se serializan a JSON y el resultado se trunca a
        `SAFE_CELL_LIMIT` (30.000) porque Excel revienta por encima de 32.767
        caracteres por celda. El truncado se marca con puntos suspensivos.
        """
        if value is None:
            return ""
        if isinstance(value, (dict, list, tuple)):
            value = json.dumps(value, ensure_ascii=False, default=str)
        else:
            value = str(value)
        if len(value) > cls.SAFE_CELL_LIMIT:
            value = value[: cls.SAFE_CELL_LIMIT] + "…"
        return value

    @classmethod
    def _local_time(cls, value):
        """Pasa un instante UTC a hora de Chile y le quita la zona horaria.

        Se quita el tzinfo porque Excel no maneja datetimes con zona; el
        contexto queda documentado en `DISPLAY_TIMEZONE`. Ante cualquier fallo
        de conversion se devuelve el valor original en lugar de perder el dato.
        """
        if value in (None, ""):
            return None
        try:
            dt = pd.to_datetime(value, utc=True, errors="coerce")
            if pd.isna(dt):
                return value
            return dt.tz_convert(ZoneInfo(cls.DISPLAY_TIMEZONE)).tz_localize(None)
        except Exception:
            return value

    @staticmethod
    def _first(*values):
        """Primer valor no vacio de la lista, o `None`.

        Sirve para el patron "tomalo de la tesis de entrada y si no del trade".
        """
        for value in values:
            if value not in (None, ""):
                return value
        return None

    def _summary_frame(self, payload):
        """Hoja "Resumen": ficha campo/valor de la operacion.

        Incluye deliberadamente el MFE/MAE persistido JUNTO al observado en los
        snapshots: si ambos difieren mucho, el seguimiento en vivo se perdio
        movimientos y esa discrepancia es en si misma un hallazgo.

        Tambien expone el bloque de integridad (origen de la tesis, snapshots
        excluidos por conflicto de identidad) y el "Motivo salida exacto",
        tomado primero de `analysis_exit_reason` para no degradarlo a una
        etiqueta generica.
        """
        trade = payload.get("trade") or {}
        snapshots = payload.get("snapshots") or []
        integrity = payload.get("audit_integrity") or {}
        entry_view = payload.get("entry_view") or (snapshots[0].get("entry_view") if snapshots else {}) or {}
        rr_values = []
        for snap in snapshots:
            try:
                rr = float((snap.get("market") or {}).get("current_rr"))
                rr_values.append(rr)
            except (TypeError, ValueError):
                pass
        rows = [
            ("Trade ID", trade.get("id")),
            ("Source Trade ID", trade.get("source_trade_id")),
            ("Instrumento", trade.get("instrument")),
            ("Estrategia ID", self._first(entry_view.get("strategy_name"), trade.get("strategy_name"))),
            ("Versión estrategia", self._first(entry_view.get("strategy_version"), trade.get("strategy_version"))),
            ("Perfil bot", self._first(entry_view.get("bot_profile"), trade.get("bot_profile"))),
            ("Magic estrategia", self._first(entry_view.get("daemon_magic"), trade.get("daemon_magic"))),
            ("Dirección", trade.get("direction")),
            ("Estado", trade.get("status")),
            ("Clasificación", trade.get("classification") or trade.get("result")),
            ("Pierna", trade.get("leg") or trade.get("execution_mode")),
            ("Entrada", self._local_time(trade.get("entry_time"))),
            ("Salida", self._local_time(trade.get("exit_time"))),
            ("Motivo salida exacto", self._first(
                (trade.get("details") or {}).get("metadata", {}).get("analysis_exit_reason"),
                trade.get("exit_reason"),
            )),
            ("RR plan", trade.get("planned_rr")),
            ("RR real", trade.get("realized_rr")),
            ("PnL neto", trade.get("net_pnl")),
            ("Riesgo %", trade.get("risk_percent")),
            ("MFE persistido", trade.get("mfe_rr")),
            ("MAE persistido", trade.get("mae_rr")),
            ("MFE observado snapshots", max(rr_values) if rr_values else None),
            ("MAE observado snapshots", min(rr_values) if rr_values else None),
            ("Snapshots", len(snapshots)),
            ("Integridad identidad", "VALIDA" if integrity.get("valid") else "INVALIDA"),
            ("Origen tesis entrada", integrity.get("entry_view_origin")),
            ("Tesis entrada completa", integrity.get("entry_view_complete")),
            ("Snapshots incompatibles excluidos", integrity.get("identity_conflicts_removed", 0)),
            ("IDs excluidos", self._safe_text(integrity.get("excluded_snapshot_ids") or [])),
            ("Generado Chile", datetime.now(ZoneInfo(self.DISPLAY_TIMEZONE)).replace(tzinfo=None)),
        ]
        return pd.DataFrame(rows, columns=["Campo", "Valor"])

    def _view_frame(self, view, title):
        """Aplana un diccionario de vista a dos columnas clave/valor.

        Emite primero las claves de `preferred` en orden fijo (decision,
        estructura, patrones, conflictos, ORB, HTF...) para que hojas de trades
        distintos sean comparables visualmente, y despues vuelca el resto de
        claves para no perder informacion.
        """
        view = view or {}
        rows = []
        preferred = [
            "decision", "state", "action", "direction", "score", "trade_score",
            "confirmation_percentage", "h1_trend", "structure_break", "zone",
            "chart_pattern_name", "chart_pattern_direction", "chart_pattern_strength",
            "chart_pattern_conflict", "chart_pattern_conflict_level",
            "chart_pattern_conflict_reason", "divergence_type", "h1_doji_type",
            "h1_doji_zone", "orb_market", "opening_range_high", "opening_range_low",
            "opening_range_midpoint", "session_vwap", "session_poc", "breakout_up",
            "breakout_down", "retest_buy_ok", "retest_sell_ok", "htf_alignment",
            "htf_blocked", "htf_reason", "evaluated_at",
        ]
        emitted = set()
        for key in preferred:
            if key in view:
                rows.append((key, self._safe_text(view.get(key))))
                emitted.add(key)
        for key, value in view.items():
            if key not in emitted:
                rows.append((key, self._safe_text(value)))
        return pd.DataFrame(rows, columns=[title, "Valor"])

    def _timeline_frame(self, payload):
        """Hoja "Timeline Completo": una fila por snapshot.

        Cruza en cada instante la tesis de entrada, la vista actual, el mercado
        y el contexto visual. Es la vista que permite datar con precision el
        momento en que la operacion dejo de comportarse como se esperaba.
        """
        rows = []
        for index, snap in enumerate(payload.get("snapshots") or [], start=1):
            entry = snap.get("entry_view") or {}
            current = snap.get("current_view") or {}
            market = snap.get("market") or {}
            visual = snap.get("visual_context") or {}
            rows.append({
                "N°": index,
                "Snapshot ID": snap.get("id"),
                "Fecha/Hora Chile": self._local_time(snap.get("snapshot_at")),
                "Bot": snap.get("bot_profile"),
                "Magic": snap.get("daemon_magic"),
                "Estrategia ID": self._first(entry.get("strategy_name"), current.get("strategy_name")),
                "Versión estrategia": entry.get("strategy_version"),
                "Instrumento": snap.get("instrument"),
                "Ticket": snap.get("broker_position_ticket"),
                "Decisión entrada": self._first(entry.get("decision"), entry.get("state"), entry.get("action")),
                "Score entrada": self._first(entry.get("score"), entry.get("trade_score")),
                "Confirmación entrada %": entry.get("confirmation_percentage"),
                "Decisión ahora": self._first(current.get("decision"), current.get("state"), current.get("action")),
                "Estado ahora": current.get("state"),
                "Análisis actual válido": current.get("valid"),
                "Motivo análisis actual": current.get("reason"),
                "Motivo salida exacto": self._first(
                    (snap.get("metadata") or {}).get("analysis_exit_reason"),
                    snap.get("exit_reason"),
                ),
                "Dirección ahora": current.get("direction"),
                "Score ahora": self._first(current.get("score"), current.get("trade_score")),
                "Confirmación ahora %": current.get("confirmation_percentage"),
                "R actual": market.get("current_rr"),
                "Precio actual": market.get("current_price"),
                "SL actual": market.get("current_stop_loss"),
                "SL inicial": market.get("initial_stop_loss"),
                "TP": market.get("take_profit"),
                "Distancia SL R": market.get("distance_to_sl_r"),
                "Distancia TP R": market.get("distance_to_tp_r"),
                "BE confirmado": market.get("break_even_confirmed"),
                "MFE R": market.get("max_favorable_excursion_rr"),
                "MAE R": market.get("max_adverse_excursion_rr"),
                "Runner etapa": market.get("runner_extension_stage"),
                "Runner lock R": market.get("runner_profit_lock_rr"),
                "H1 ahora": current.get("h1_trend"),
                "Estructura ahora": current.get("structure_break"),
                "Patrón chartista": current.get("chart_pattern_name"),
                "Fuerza patrón": current.get("chart_pattern_strength"),
                "Conflicto chartista": current.get("chart_pattern_conflict"),
                "Nivel conflicto": current.get("chart_pattern_conflict_level"),
                "Motivo conflicto": current.get("chart_pattern_conflict_reason"),
                "Divergencia": current.get("divergence_type"),
                "Doji H1": current.get("h1_doji_type"),
                "ORB mercado": current.get("orb_market"),
                "ORB high": current.get("opening_range_high"),
                "ORB low": current.get("opening_range_low"),
                "ORB midpoint": current.get("opening_range_midpoint"),
                "VWAP": current.get("session_vwap"),
                "POC": current.get("session_poc"),
                "HTF alineado": current.get("htf_alignment"),
                "HTF bloqueado": current.get("htf_blocked"),
                "HTF motivo": current.get("htf_reason"),
                "Confirmaciones aprobadas": self._safe_text(current.get("passed_confirmations") or current.get("confirmations_passed")),
                "Confirmaciones faltantes": self._safe_text(current.get("missing_confirmations")),
                "Entry JSON": self._safe_text(entry),
                "Current JSON": self._safe_text(current),
                "Market JSON": self._safe_text(market),
                "Visual Context JSON": self._safe_text(visual),
            })
        return pd.DataFrame(rows)

    def export(self, payload):
        """Valida la integridad y escribe el libro completo de auditoria.

        Args:
            payload: dict con `trade`, `snapshots`, `entry_view`, `latest` y
                `audit_integrity`.

        Returns:
            dict con `path`, `trade_id`, `instrument` y numero de `snapshots`.

        Raises:
            ValueError: `AUDIT_INTEGRITY_ERROR` si la integridad viene marcada
                como invalida, si falta la tesis de entrada o si algun snapshot
                pertenece a otro instrumento. Se falla en vez de exportar un
                informe forense potencialmente enganoso.
        """
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        trade = payload.get("trade") or {}
        snapshots = payload.get("snapshots") or []
        entry_view = payload.get("entry_view") or (snapshots[0].get("entry_view") if snapshots else {}) or {}
        integrity = payload.get("audit_integrity") or {}
        if integrity and not integrity.get("valid", False):
            raise ValueError("AUDIT_INTEGRITY_ERROR: identidad del trade no verificable")
        expected_instrument = str(trade.get("instrument") or "")
        mismatches = [
            snap.get("id") for snap in snapshots
            if str(snap.get("instrument") or "") != expected_instrument
        ]
        if not entry_view or mismatches:
            raise ValueError(
                "AUDIT_INTEGRITY_ERROR: tesis ausente o instrumentos mezclados "
                f"(snapshots={mismatches})"
            )
        latest = payload.get("latest") or (snapshots[-1] if snapshots else {}) or {}
        current_view = latest.get("current_view") or {}
        latest_market = latest.get("market") or {}

        summary_df = self._summary_frame(payload)
        entry_df = self._view_frame(entry_view, "Tesis de entrada")
        current_df = self._view_frame(current_view, "Último estado")
        market_df = self._view_frame(latest_market, "Mercado final")
        timeline_df = self._timeline_frame(payload)

        with pd.ExcelWriter(self.output_path, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="Resumen", index=False)
            entry_df.to_excel(writer, sheet_name="Tesis Entrada", index=False)
            current_df.to_excel(writer, sheet_name="Ultimo Estado", index=False)
            market_df.to_excel(writer, sheet_name="Mercado Final", index=False)
            timeline_df.to_excel(writer, sheet_name="Timeline Completo", index=False)

        self._format_workbook()
        return {
            "path": str(self.output_path),
            "trade_id": trade.get("id"),
            "instrument": trade.get("instrument"),
            "snapshots": len(snapshots),
        }

    def _format_workbook(self):
        """Aplica estilo al libro ya escrito: cabecera, filtros y anchos.

        Congela la fila de titulos, activa autofiltro y ajusta el ancho de
        columna inspeccionando solo las 200 primeras filas (medir un timeline
        largo entero seria costoso y no cambia el resultado). El ancho se acota
        entre 10 y 45 para que las celdas JSON no desborden la pantalla.

        `openpyxl` se importa aqui dentro para no pagar el coste al importar el
        modulo cuando no se va a exportar nada.
        """
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = load_workbook(self.output_path)
        header_fill = PatternFill("solid", fgColor="1B1F23")
        header_font = Font(color="F5C84B", bold=True)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column_cells in ws.columns:
                letter = column_cells[0].column_letter
                max_len = 0
                for cell in column_cells[:200]:
                    value = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, min(len(value), 45))
                ws.column_dimensions[letter].width = max(10, min(max_len + 2, 45))
        wb.save(self.output_path)
